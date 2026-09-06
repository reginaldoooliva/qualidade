import statistics
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from app.models.nao_conformidade import StatusNC
from app.models.peca import StatusCadastro
from app.repositories import (
    analise_repository,
    caracteristica_repository,
    etapa_repository,
    nao_conformidade_repository,
    peca_repository,
    plano_acao_repository,
    relatorio_repository,
)
from app.services import capabilidade_service, nao_conformidade_service, peca_service, plano_acao_service
from app.services.relatorio_pdf import gerar_pdf_consolidado_nc, gerar_pdf_consolidado_pecas, gerar_pdf_cpk_peca, gerar_pdf_rnc


_CLASSIFICACAO_NC = {"critica": "Crítica", "maior": "Maior", "menor": "Menor"}
_ORIGEM_NC = {
    "processo": "Processo",
    "materia_prima": "Matéria-prima",
    "projeto": "Projeto/Desenho",
    "instrumento": "Instrumento de medição",
    "mao_de_obra": "Mão de obra",
    "outro": "Outro",
}
_DISPOSICAO_NC = {
    "retrabalho": "Retrabalho",
    "sucata": "Sucata",
    "uso_como_esta": "Uso como está (concessão)",
    "devolucao_fornecedor": "Devolução ao fornecedor",
    "reclassificacao": "Reclassificação",
}
_STATUS_NC = {"aberta": "Aberta", "em_analise": "Em análise", "em_tratamento": "Em tratamento", "encerrada": "Encerrada"}
_METODOLOGIA = {"ishikawa": "Ishikawa (Espinha de Peixe)", "cinco_porques": "5 Porquês", "livre": "Análise livre"}
_STATUS_PLANO = {
    "aberto": "Aberto",
    "em_andamento": "Em andamento",
    "aguardando_verificacao": "Aguardando verificação",
    "encerrado": "Encerrado",
    "reaberto": "Reaberto",
}
_STATUS_ACAO = {"pendente": "Pendente", "concluida": "Concluída", "atrasada": "Atrasada"}
_RESULTADO_VERIFICACAO = {"eficaz": "Eficaz", "nao_eficaz": "Não eficaz"}


def _filtro_texto(data_inicio: date | None, data_fim: date | None, operador_nome: str | None) -> str | None:
    partes = []
    if data_inicio or data_fim:
        partes.append(f"Período: {data_inicio or '...'} a {data_fim or '...'}")
    if operador_nome:
        partes.append(f"Operador: {operador_nome}")
    return " · ".join(partes) if partes else None


def cpk_pdf_peca(
    db: Session,
    peca_id: int,
    data_inicio: date | None = None,
    data_fim: date | None = None,
    operador_id: int | None = None,
) -> bytes:
    peca = peca_service.obter(db, peca_id)
    etapas = [e for e in etapa_repository.list_by_peca(db, peca_id) if e.status == StatusCadastro.ATIVO]

    operador_nome = None
    etapas_dados = []
    for etapa in etapas:
        resultado = capabilidade_service.analisar(
            db, peca_id, etapa.id, data_inicio=data_inicio, data_fim=data_fim, operador_id=operador_id
        )
        if operador_id is not None:
            encontrado = next((o for o in resultado["operadores_disponiveis"] if o.id == operador_id), None)
            operador_nome = encontrado.nome if encontrado else operador_nome
        etapas_dados.append(
            {
                "numero_etapa": etapa.numero_etapa,
                "descricao": etapa.descricao,
                "caracteristicas": resultado["caracteristicas"],
            }
        )

    filtro_texto = _filtro_texto(data_inicio, data_fim, operador_nome)
    return gerar_pdf_cpk_peca(
        {"codigo": peca.codigo, "descricao": peca.descricao}, etapas_dados, filtro_texto
    )


def dados_brutos_xlsx_peca(
    db: Session,
    peca_id: int,
    etapa_id: int | None = None,
    data_inicio: date | None = None,
    data_fim: date | None = None,
    operador_id: int | None = None,
) -> bytes:
    peca = peca_service.obter(db, peca_id)
    medicoes = relatorio_repository.list_medicoes_peca(
        db, peca_id, etapa_id=etapa_id, data_inicio=data_inicio, data_fim=data_fim, operador_id=operador_id
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Dados brutos"
    cabecalho = [
        "Ordem",
        "Etapa",
        "Característica",
        "Amostra",
        "Valor",
        "Operador",
        "Data/hora",
    ]
    ws.append(cabecalho)
    for celula in ws[1]:
        celula.font = Font(bold=True)

    for m in medicoes:
        ws.append(
            [
                m.rodada.ordem.numero_ordem,
                f"Etapa {m.rodada.etapa.numero_etapa}",
                m.caracteristica.nome,
                m.amostra_numero,
                float(m.valor),
                m.operador.nome,
                m.data_hora.strftime("%d/%m/%Y %H:%M"),
            ]
        )

    larguras = [14, 12, 24, 10, 12, 20, 18]
    for idx, largura in enumerate(larguras, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = largura

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def consolidado_pecas(db: Session) -> list[dict]:
    pecas = [p for p in peca_repository.list_pecas(db) if p.status == StatusCadastro.ATIVO]

    resultados = []
    for peca in pecas:
        pior: dict | None = None
        etapas = [e for e in etapa_repository.list_by_peca(db, peca.id) if e.status == StatusCadastro.ATIVO]
        for etapa in etapas:
            rodadas = analise_repository.list_rodadas_para_analise(db, peca.id, etapa.id)
            rodada_ids = [r.id for r in rodadas]
            caracteristicas = [
                c for c in caracteristica_repository.list_by_etapa(db, etapa.id) if c.status == StatusCadastro.ATIVO
            ]
            for caracteristica in caracteristicas:
                medicoes = analise_repository.list_medicoes(db, rodada_ids, caracteristica.id)
                valores = [float(m.valor) for m in medicoes]
                stats = capabilidade_service.calcular_estatisticas(valores, caracteristica.lie, caracteristica.lse)
                if stats["cpk"] is None:
                    continue
                if pior is None or stats["cpk"] < pior["cpk_critico"]:
                    pior = {
                        "caracteristica_critica": caracteristica.nome,
                        "etapa_numero": etapa.numero_etapa,
                        "cpk_critico": stats["cpk"],
                        "classificacao": stats["classificacao"],
                        "n_amostras": stats["n_amostras"],
                    }

        resultados.append(
            {
                "peca_id": peca.id,
                "codigo": peca.codigo,
                "descricao": peca.descricao,
                "caracteristica_critica": pior["caracteristica_critica"] if pior else None,
                "etapa_numero": pior["etapa_numero"] if pior else None,
                "cpk_critico": pior["cpk_critico"] if pior else None,
                "classificacao": pior["classificacao"] if pior else None,
                "n_amostras": pior["n_amostras"] if pior else 0,
            }
        )

    resultados.sort(key=lambda r: (r["cpk_critico"] is None, r["cpk_critico"]))
    return resultados


def consolidado_pecas_pdf(db: Session) -> bytes:
    return gerar_pdf_consolidado_pecas(consolidado_pecas(db))


def consolidado_nao_conformidades(db: Session, data_inicio: date | None = None, data_fim: date | None = None) -> dict:
    ncs = nao_conformidade_repository.list_nc(db, data_inicio=data_inicio, data_fim=data_fim)

    por_peca: dict[str, int] = {}
    por_classificacao: dict[str, int] = {}
    por_origem: dict[str, int] = {}
    tempos_tratamento: list[float] = []
    total_encerradas = 0

    for nc in ncs:
        chave_peca = f"{nc.peca.codigo} — {nc.peca.descricao}"
        por_peca[chave_peca] = por_peca.get(chave_peca, 0) + 1
        por_classificacao[nc.classificacao.value] = por_classificacao.get(nc.classificacao.value, 0) + 1
        por_origem[nc.origem.value] = por_origem.get(nc.origem.value, 0) + 1
        if nc.status == StatusNC.ENCERRADA and nc.data_encerramento is not None:
            total_encerradas += 1
            dias = (nc.data_encerramento - nc.data_abertura).total_seconds() / 86400
            tempos_tratamento.append(dias)

    planos_todos = [plano for nc in ncs for plano in nc.planos_acao]
    reabertos = sum(1 for plano in planos_todos if plano.ciclo > 1)
    taxa_reincidencia = reabertos / len(planos_todos) if planos_todos else 0.0

    return {
        "total_rnc": len(ncs),
        "total_encerradas": total_encerradas,
        "tempo_medio_tratamento_dias": statistics.fmean(tempos_tratamento) if tempos_tratamento else None,
        "taxa_reincidencia": taxa_reincidencia,
        "por_peca": por_peca,
        "por_classificacao": por_classificacao,
        "por_origem": por_origem,
    }


def consolidado_nao_conformidades_pdf(db: Session, data_inicio: date | None = None, data_fim: date | None = None) -> bytes:
    dados = consolidado_nao_conformidades(db, data_inicio=data_inicio, data_fim=data_fim)
    dados["por_classificacao"] = {_CLASSIFICACAO_NC[k]: v for k, v in dados["por_classificacao"].items()}
    dados["por_origem"] = {_ORIGEM_NC[k]: v for k, v in dados["por_origem"].items()}
    return gerar_pdf_consolidado_nc(dados)


def rnc_pdf(db: Session, nc_id: int) -> bytes:
    nc = nao_conformidade_service.obter(db, nc_id)
    detalhe = nao_conformidade_service.to_detalhe(db, nc)

    nc_dict = {
        "numero_rnc": detalhe.numero_rnc,
        "peca_codigo": detalhe.peca.codigo,
        "peca_descricao": detalhe.peca.descricao,
        "descricao_problema": detalhe.descricao_problema,
        "quantidade_afetada": detalhe.quantidade_afetada,
        "classificacao": _CLASSIFICACAO_NC[detalhe.classificacao.value],
        "origem": _ORIGEM_NC[detalhe.origem.value],
        "status": _STATUS_NC[detalhe.status.value],
        "aberto_por_nome": detalhe.aberto_por_nome,
        "data_abertura": detalhe.data_abertura.strftime("%d/%m/%Y %H:%M"),
        "causa_raiz_preliminar": detalhe.causa_raiz_preliminar,
        "disposicao": _DISPOSICAO_NC[detalhe.disposicao.value] if detalhe.disposicao else None,
        "historico": [
            {
                "criado_em": evento.criado_em.strftime("%d/%m/%Y %H:%M"),
                "usuario_nome": evento.usuario_nome,
                "detalhe": evento.detalhe,
                "acao": evento.acao,
            }
            for evento in detalhe.historico
        ],
    }

    planos_dict = []
    for resumo in detalhe.planos_acao:
        plano = plano_acao_service.obter(db, resumo.id)
        plano_detalhe = plano_acao_service.to_detalhe(db, plano)
        planos_dict.append(
            {
                "id": plano_detalhe.id,
                "ciclo": plano_detalhe.ciclo,
                "status": _STATUS_PLANO[plano_detalhe.status.value],
                "metodologia_causa_raiz": _METODOLOGIA[plano_detalhe.metodologia_causa_raiz.value],
                "conclusao_causa_raiz": plano_detalhe.conclusao_causa_raiz,
                "acoes_corretivas": [
                    {
                        "descricao": a.descricao,
                        "responsavel_nome": a.responsavel_nome,
                        "prazo": a.prazo.strftime("%d/%m/%Y"),
                        "status": _STATUS_ACAO[a.status],
                    }
                    for a in plano_detalhe.acoes_corretivas
                ],
                "verificacoes": [
                    {
                        "data_verificacao": v.data_verificacao.strftime("%d/%m/%Y"),
                        "responsavel_nome": v.responsavel_nome,
                        "resultado": _RESULTADO_VERIFICACAO[v.resultado.value],
                        "observacoes": v.observacoes,
                    }
                    for v in plano_detalhe.verificacoes
                ],
            }
        )

    return gerar_pdf_rnc(nc_dict, planos_dict)
